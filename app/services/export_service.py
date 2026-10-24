import io
import csv
from datetime import datetime
from pathlib import Path
from typing import List

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
from PIL import Image
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.config import settings
from app.models.report import Report

class ExportService:

    @staticmethod
    def generate_report_pdf(report: Report) -> io.BytesIO:
        """Genera un informe técnico en PDF de alta calidad estética con ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Estilos personalizados
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#0f172a'),
            alignment=TA_LEFT
        )
        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=13,
            textColor=colors.HexColor('#64748b'),
            alignment=TA_LEFT
        )
        section_style = ParagraphStyle(
            'ReportSection',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=15,
            textColor=colors.HexColor('#1e293b'),
            spaceBefore=8,
            spaceAfter=4
        )
        body_style = ParagraphStyle(
            'ReportBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#334155'),
            alignment=TA_JUSTIFY
        )
        badge_style = ParagraphStyle(
            'BadgeStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=12,
            textColor=colors.HexColor('#ffffff'),
            alignment=TA_CENTER
        )

        elements = []

        # 1. ENCABEZADO INSTITUCIONAL
        header_data = [
            [
                Paragraph("<b>DEPARTAMENTO DE HIGIENE Y SEGURIDAD EN EL TRABAJO</b><br/><font size=8 color='#64748b'>Taller Metalúrgico y Soldadura Industrial S.A.</font>", title_style),
                Paragraph(f"<font size=14 color='#0284c7'><b>{report.code}</b></font><br/><font size=8 color='#64748b'>Fecha: {report.created_at.strftime('%d/%m/%Y %H:%M')}</font>", ParagraphStyle('HRight', parent=subtitle_style, alignment=TA_RIGHT))
            ]
        ]
        header_table = Table(header_data, colWidths=[380, 160])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(header_table)
        elements.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#0284c7'), spaceAfter=10))

        # 2. TABLA DE METADATOS PRINCIPALES
        status_colors = {
            "Reportado": "#3b82f6",
            "En revisión": "#6366f1",
            "En análisis": "#8b5cf6",
            "Acción asignada": "#0284c7",
            "En proceso": "#f59e0b",
            "Pendiente de verificación": "#f97316",
            "Cerrado": "#10b981",
            "Descartado": "#64748b"
        }
        st_color = status_colors.get(report.status, "#475569")

        reporter_display = "Anónimo" if report.is_anonymous else (report.reporter.full_name if report.reporter else "No registrado")
        
        info_data = [
            [
                Paragraph("<b>Tipo de Evento:</b>", body_style),
                Paragraph(f"{report.report_type.name}", body_style),
                Paragraph("<b>Estado Actual:</b>", body_style),
                Paragraph(f"<font color='{st_color}'><b>{report.status.upper()}</b></font>", body_style)
            ],
            [
                Paragraph("<b>Ubicación / Sector:</b>", body_style),
                Paragraph(f"{report.location.name}", body_style),
                Paragraph("<b>Nivel de Urgencia:</b>", body_style),
                Paragraph(f"{report.urgency}", body_style)
            ],
            [
                Paragraph("<b>Reportado por:</b>", body_style),
                Paragraph(reporter_display, body_style),
                Paragraph("<b>Peligro Inmediato:</b>", body_style),
                Paragraph("<font color='red'><b>SÍ (Crítico)</b></font>" if report.immediate_danger else "NO", body_style)
            ]
        ]
        info_table = Table(info_data, colWidths=[110, 160, 120, 150])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 10))

        # 3. DESCRIPCIÓN DEL HECHO
        elements.append(Paragraph("1. DESCRIPCIÓN DE LA SITUACIÓN OBSERVADA", section_style))
        desc_box = Table([[Paragraph(report.description, body_style)]], colWidths=[540])
        desc_box.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#ffffff')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ]))
        elements.append(desc_box)
        elements.append(Spacer(1, 10))

        # 4. EVALUACIÓN TÉCNICA DEL RIESGO
        elements.append(Paragraph("2. EVALUACIÓN TÉCNICA DE RIESGO (MATRIZ 5x5)", section_style))
        if report.risk_assessment:
            ra = report.risk_assessment
            risk_colors = {
                "Bajo": "#10b981",
                "Moderado": "#eab308",
                "Importante": "#f97316",
                "Alto": "#ef4444",
                "Crítico": "#991b1b"
            }
            rc = risk_colors.get(ra.risk_level, "#334155")

            risk_table_data = [
                [
                    Paragraph("<b>Probabilidad (P)</b>", body_style),
                    Paragraph("<b>Consecuencia (C)</b>", body_style),
                    Paragraph("<b>Puntaje (R = P x C)</b>", body_style),
                    Paragraph("<b>Nivel de Riesgo</b>", body_style)
                ],
                [
                    Paragraph(f"{ra.probability} / 5", body_style),
                    Paragraph(f"{ra.consequence} / 5", body_style),
                    Paragraph(f"<b>{ra.risk_score}</b> / 25", body_style),
                    Paragraph(f"<font color='{rc}'><b>{ra.risk_level.upper()}</b></font>", body_style)
                ]
            ]
            if ra.observations:
                risk_table_data.append([
                    Paragraph("<b>Observaciones Técnicas:</b>", body_style),
                    Paragraph(ra.observations, body_style),
                    Paragraph("", body_style),
                    Paragraph("", body_style)
                ])

            risk_table = Table(risk_table_data, colWidths=[120, 120, 140, 160])
            risk_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
                ('ALIGN', (0, 0), (-1, 1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('SPAN', (1, 2), (3, 2)) if ra.observations else ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ]))
            elements.append(risk_table)
        else:
            elements.append(Paragraph("<i>Pendiente de evaluación técnica por el Responsable de Seguridad.</i>", subtitle_style))
        elements.append(Spacer(1, 10))

        # 5. PLAN DE ACCIONES CORRECTIVAS Y PREVENTIVAS
        elements.append(Paragraph("3. PLAN DE ACCIONES Y JERARQUÍA DE CONTROLES", section_style))
        if report.actions:
            actions_data = [
                [
                    Paragraph("<b>#</b>", body_style),
                    Paragraph("<b>Tipo / Control</b>", body_style),
                    Paragraph("<b>Descripción de la Acción</b>", body_style),
                    Paragraph("<b>Responsable</b>", body_style),
                    Paragraph("<b>Vencimiento</b>", body_style),
                    Paragraph("<b>Estado</b>", body_style)
                ]
            ]
            for idx, a in enumerate(report.actions, 1):
                resp = a.assigned_to.full_name if a.assigned_to else "Sin asignar"
                due = a.due_date.strftime("%d/%m/%Y") if a.due_date else "-"
                actions_data.append([
                    Paragraph(str(idx), body_style),
                    Paragraph(f"<b>{a.action_type}</b><br/><font size=7 color='#64748b'>{a.control_hierarchy}</font>", body_style),
                    Paragraph(a.description, body_style),
                    Paragraph(resp, body_style),
                    Paragraph(due, body_style),
                    Paragraph(f"<b>{a.status}</b>", body_style)
                ])

            actions_table = Table(actions_data, colWidths=[20, 110, 190, 85, 65, 70])
            actions_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            elements.append(actions_table)
        else:
            elements.append(Paragraph("<i>No se han asignado acciones hasta el momento.</i>", subtitle_style))
        elements.append(Spacer(1, 10))

        # 6. EVIDENCIA FOTOGRÁFICA
        report_photos = [att for att in report.attachments if not att.is_evidence]
        evidence_photos = [att for att in report.attachments if att.is_evidence]

        if report_photos or evidence_photos:
            elements.append(Paragraph("4. REGISTRO FOTOGRÁFICO Y EVIDENCIAS", section_style))
            photo_cells = []
            
            # Incorporar foto inicial si existe
            if report_photos:
                first_photo = report_photos[0]
                photo_name = Path(first_photo.file_path).name
                local_path = settings.UPLOAD_DIR / photo_name
                if not local_path.exists():
                    local_path = settings.BASE_DIR / first_photo.file_path.lstrip("/")
                if local_path.exists():
                    try:
                        with Image.open(local_path) as test_im:
                            test_im.verify()
                        img = RLImage(str(local_path), width=240, height=160)
                        photo_cells.append([img, Paragraph("<b>Foto inicial del reporte:</b><br/>Condición/acto observado en planta.", body_style)])
                    except Exception:
                        pass

            if evidence_photos:
                first_ev = evidence_photos[0]
                ev_name = Path(first_ev.file_path).name
                local_ev_path = settings.UPLOAD_DIR / ev_name
                if not local_ev_path.exists():
                    local_ev_path = settings.BASE_DIR / first_ev.file_path.lstrip("/")
                if local_ev_path.exists():
                    try:
                        with Image.open(local_ev_path) as test_im:
                            test_im.verify()
                        img_ev = RLImage(str(local_ev_path), width=240, height=160)
                        photo_cells.append([img_ev, Paragraph("<b>Evidencia de resolución:</b><br/>Medida de control implementada.", body_style)])
                    except Exception:
                        pass

            if photo_cells:
                photo_table = Table(photo_cells, colWidths=[260, 280])
                photo_table.setStyle(TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ]))
                elements.append(KeepTogether(photo_table))
            elements.append(Spacer(1, 10))

        # 7. CIERRE Y FIRMAS FORMALES
        elements.append(Spacer(1, 15))
        if report.status == "Cerrado":
            close_text = f"<b>Fecha de Cierre:</b> {report.closed_at.strftime('%d/%m/%Y %H:%M') if report.closed_at else '-'}<br/><b>Dictamen de Cierre:</b> {report.close_reason or 'Verificado satisfactoriamente.'}"
            elements.append(Table([[Paragraph(close_text, body_style)]], colWidths=[540], style=[
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#ecfdf5')),
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#10b981')),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            elements.append(Spacer(1, 20))

        sign_data = [
            [
                Paragraph("________________________________________<br/><b>Responsable de Seguridad y Salud (SST)</b><br/>Firma y Sello Profesional", ParagraphStyle('Sign1', parent=body_style, alignment=TA_CENTER)),
                Paragraph("________________________________________<br/><b>Jefatura de Planta / Mantenimiento</b><br/>Firma de Conformidad", ParagraphStyle('Sign2', parent=body_style, alignment=TA_CENTER))
            ]
        ]
        sign_table = Table(sign_data, colWidths=[270, 270])
        sign_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 15),
        ]))
        elements.append(KeepTogether(sign_table))

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_excel_export(reports: List[Report]) -> io.BytesIO:
        """Genera una planilla de cálculo Excel (.xlsx) estilizada profesionalmente."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reportes SST"

        # Encabezado corporativo
        ws.merge_cells("A1:M1")
        header_cell = ws["A1"]
        header_cell.value = "SISTEMA DE GESTIÓN DE SEGURIDAD Y SALUD EN EL TRABAJO - REGISTRO DE REPORTES"
        header_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        header_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Subtítulo con fecha
        ws.merge_cells("A2:M2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Exportación generada el {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} - Taller Metalúrgico de Soldadura"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        columns = [
            ("Código", 15),
            ("Fecha", 14),
            ("Tipo", 22),
            ("Ubicación", 24),
            ("Descripción", 45),
            ("Urgencia", 12),
            ("Peligro Inmediato", 16),
            ("Reportante", 22),
            ("Estado", 18),
            ("Puntaje Riesgo", 14),
            ("Nivel Riesgo", 14),
            ("Acciones Asignadas", 16),
            ("Fecha Cierre", 14),
        ]

        # Fila de encabezados de columnas
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for col_idx, (col_name, col_width) in enumerate(columns, 1):
            cell = ws.cell(row=3, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_width

        ws.row_dimensions[3].height = 25

        # Datos
        for row_idx, r in enumerate(reports, 4):
            reporter_str = "Anónimo" if r.is_anonymous else (r.reporter.full_name if r.reporter else "No registrado")
            risk_sc = r.risk_assessment.risk_score if r.risk_assessment else "-"
            risk_lvl = r.risk_assessment.risk_level if r.risk_assessment else "Sin evaluar"
            close_dt = r.closed_at.strftime("%d/%m/%Y") if r.closed_at else "-"

            row_data = [
                r.code,
                r.created_at.strftime("%d/%m/%Y"),
                r.report_type.name,
                r.location.name,
                r.description,
                r.urgency,
                "SÍ" if r.immediate_danger else "NO",
                reporter_str,
                r.status,
                risk_sc,
                risk_lvl,
                len(r.actions),
                close_dt
            ]

            fill_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                
                # Alineación
                if col_idx in [1, 2, 6, 7, 9, 10, 11, 12, 13]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[row_idx].height = 20

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_csv_export(reports: List[Report]) -> io.StringIO:
        """Genera un archivo CSV estándar con codificación UTF-8."""
        buffer = io.StringIO()
        # Escribir BOM para Excel
        buffer.write('\ufeff')
        writer = csv.writer(buffer, delimiter=';')

        writer.writerow([
            "Codigo", "Fecha", "Tipo", "Ubicacion", "Descripcion", "Urgencia",
            "PeligroInmediato", "Reportante", "Estado", "PuntajeRiesgo",
            "NivelRiesgo", "CantidadAcciones", "FechaCierre"
        ])

        for r in reports:
            reporter_str = "Anonimo" if r.is_anonymous else (r.reporter.full_name if r.reporter else "No registrado")
            risk_sc = r.risk_assessment.risk_score if r.risk_assessment else ""
            risk_lvl = r.risk_assessment.risk_level if r.risk_assessment else "Sin evaluar"
            close_dt = r.closed_at.strftime("%Y-%m-%d") if r.closed_at else ""

            writer.writerow([
                r.code,
                r.created_at.strftime("%Y-%m-%d %H:%M"),
                r.report_type.name,
                r.location.name,
                r.description.replace("\n", " "),
                r.urgency,
                "SI" if r.immediate_danger else "NO",
                reporter_str,
                r.status,
                risk_sc,
                risk_lvl,
                len(r.actions),
                close_dt
            ])

        buffer.seek(0)
        return buffer

export_service = ExportService()
